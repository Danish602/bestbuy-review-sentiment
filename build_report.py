"""Builds the Word report from the CSV the notebook produced.

Usage:
    python build_report.py                                   # real run  -> data/bestbuy_reviews_sentiment.csv
    python build_report.py data/demo/bestbuy_reviews_sentiment.csv Report_DEMO.docx   # demo run

Every number in the report is computed from the CSV - nothing is typed in by hand.
"""
import ast, re, sys
from collections import Counter
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Inches, RGBColor

csv_path = Path(sys.argv[1] if len(sys.argv) > 1 else "data/bestbuy_reviews_sentiment.csv")
out_path = Path(sys.argv[2] if len(sys.argv) > 2 else "Report_BestBuy_Review_Analysis.docx")
if not csv_path.exists():
    sys.exit(f"Cannot find {csv_path}. Run the notebook first (or pass the CSV path).")

df = pd.read_csv(csv_path)
is_demo = df["Source"].astype(str).str.contains("SYNTHETIC", case=False).any()
chart_dir = csv_path.parent

# ---------- compute everything from the data ----------
TAG = {"Pos": "Positive", "Neg": "Negative", "Neu": "Neutral"}
counts = Counter()
for cell in df["Aspect_Sentiment"].fillna("[]"):
    try:
        items = ast.literal_eval(cell)
    except (ValueError, SyntaxError):
        continue
    for item in items:
        m = re.match(r"(.+) \((Pos|Neg|Neu)\)$", item)
        if m:
            counts[(m.group(1), TAG[m.group(2)])] += 1
aspects = sorted({a for a, _ in counts})
rows = []
for a in aspects:
    p, n, u = counts[(a, "Positive")], counts[(a, "Negative")], counts[(a, "Neutral")]
    rows.append({"Aspect": a, "Positive": p, "Negative": n, "Neutral": u, "Mentions": p + n + u})
asp = pd.DataFrame(rows)
generic = {"Overall Satisfaction", "General Satisfaction"}
drv = asp[~asp["Aspect"].isin(generic)] if not asp.empty else asp
top_pos = drv[drv["Positive"] > 0].sort_values("Positive", ascending=False).head(3) if not drv.empty else drv
top_neg = drv[drv["Negative"] > 0].sort_values("Negative", ascending=False).head(3) if not drv.empty else drv

n = len(df)
sent = df["Sentiment_Label"].value_counts()
pct = lambda k: round(100 * sent.get(k, 0) / n, 1)
avg = round(pd.to_numeric(df["Rating"], errors="coerce").mean(), 2)
dates = df["Date"].dropna().astype(str)
dates = dates[dates.str.len() == 10]
date_range = f"{dates.min()} to {dates.max()}" if len(dates) else "n/a"
mismatch = int((df["Rating_vs_Text_Flag"].fillna("") != "").sum())

# What to do about each pain point (generic, product-agnostic playbook)
RECO = {
    "Design": "Review the physical design feedback (size, look, build) and reflect it in the next product revision and in product-page photos/dimensions.",
    "Quality & Durability": "Investigate failure reports, tighten supplier quality checks, and consider a clearer warranty message; flag repeated defect themes to the vendor.",
    "Ease of Use": "Simplify onboarding: clearer in-box quick-start guide, setup video, and a short troubleshooting FAQ on the product page.",
    "Value for Money": "Test promotions or bundles, and make the value proposition clearer (what is included, comparisons) on the product page.",
    "Performance": "Share performance/lag/crash themes with the manufacturer and publish known fixes (firmware/app updates) on the product page.",
    "Display & Sound": "Set accurate expectations (screen size, audio capability) in the listing and escalate recurring complaints to the vendor.",
    "Connectivity": "Publish a connectivity/Wi-Fi troubleshooting guide and push the vendor for firmware fixes for recurring disconnect issues.",
    "Battery & Power": "Clarify battery life and power requirements in the listing; escalate recurring power issues to the vendor.",
    "Customer Service & Delivery": "Improve post-purchase experience: delivery-damage packaging checks, faster return handling, and proactive status updates.",
    "Features": "Make feature limitations explicit in the listing and collect requested features for the vendor roadmap.",
}

doc = Document()
st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(11)

def para(text, bold=False, italic=False, color=None, size=None):
    p = doc.add_paragraph(); r = p.add_run(text); r.bold, r.italic = bold, italic
    if color: r.font.color.rgb = RGBColor(*color)
    if size: r.font.size = Pt(size)
    return p
def bullet(text): doc.add_paragraph(text, style="List Bullet")
def table(header, body, widths=None):
    t = doc.add_table(rows=1, cols=len(header)); t.style = "Light Grid Accent 1"
    for i, h in enumerate(header):
        t.rows[0].cells[i].text = h
        for r in t.rows[0].cells[i].paragraphs[0].runs: r.bold = True
    for row in body:
        cells = t.add_row().cells
        for i, v in enumerate(row): cells[i].text = str(v)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths): row.cells[i].width = Inches(w)
    doc.add_paragraph()

doc.add_heading("BestBuy.ca Product Reviews - Scraping Approach and Customer Insights", 0)
if is_demo:
    para("DEMO REPORT - generated from SYNTHETIC sample data to show the format. Not real customer data. "
         "Run the notebook with USE_SAMPLE_DATA = False and re-run build_report.py to produce the real report.",
         bold=True, color=(0xC0, 0x00, 0x00))

doc.add_heading("1. Scope", 1)
para(f"{n} unique reviews were analysed ({date_range}). Average star rating: {avg} / 5. "
     "Reviews were collected from the public product page only, using Selenium with the page's own filters "
     "(Most Helpful, Newest, Highest Rating, Lowest Rating, Most Relevant) and \"Show more\" paging, then de-duplicated. "
     "Sentiment was scored with VADER (rule-based) at review level and at aspect level (design, quality, ease of use, value, etc.).")

doc.add_heading("2. Solutions for scraping challenges", 1)
table(["Challenge", "Solution implemented / recommended"], [
    ["Anti-scraping blockers & IP blocking",
     "Implemented: real Chrome browser, realistic user-agent, random delays (1.5-3.5 s), single sequential session, rest between filters. "
     "If scaling: rotating residential proxies (one IP per session), lower request rate, off-peak runs."],
    ["CAPTCHA",
     "Implemented: detection of CAPTCHA / 'access denied' pages; the scraper stops cleanly and keeps data collected so far. "
     "Recommended: do not bypass CAPTCHAs - slow down and retry later, ask the retailer for a data feed/partner API, or use a licensed review-data provider."],
    ["Network errors / timeouts",
     "Implemented: retry with exponential backoff and jitter (2 s, 4 s, 8 s), per-filter isolation, partial results saved after every filter."],
    ["Dynamic (JavaScript) content & 'Show more'",
     "Implemented: page scrolling for lazy loading; clicks 'Show more' until the review count stops growing (judged by effect, so unrelated 'show more' buttons are skipped); expands 'Read more'."],
    ["Layout / class-name changes",
     "Implemented: content-based parser anchored on review dates instead of fragile CSS classes; page HTML saved to debug/ when nothing is parsed."],
    ["Duplicate reviews across filters", "Implemented: content fingerprint used as primary key fallback; duplicates removed."],
    ["Legal / ethical",
     "Public data only, no login, polite request rate. Review the site's Terms of Use and robots.txt before any large-scale or commercial use."],
], widths=[2.0, 4.5])

doc.add_heading("3. Sentiment results", 1)
table(["Sentiment", "Reviews", "Share"], [[k, int(sent.get(k, 0)), f"{pct(k)}%"] for k in ("Positive", "Neutral", "Negative")], widths=[2, 1.5, 1.5])
for name in ("chart_sentiment.png", "chart_aspects.png"):
    if (chart_dir / name).exists():
        doc.add_picture(str(chart_dir / name), width=Inches(5.2))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
para(f"Reviews where the star rating and the text disagree: {mismatch} (flagged in the data for manual checking).", italic=True)

doc.add_heading("4. Summary of insights for stakeholders", 1)
doc.add_heading("What customers like (drivers of satisfaction)", 2)
if top_pos.empty: para("No clear positive aspect themes were detected.")
for _, r in top_pos.iterrows():
    bullet(f"{r['Aspect']}: {int(r['Positive'])} positive vs {int(r['Negative'])} negative mentions.")
doc.add_heading("What frustrates customers (drivers of dissatisfaction)", 2)
if top_neg.empty: para("No clear negative aspect themes were detected.")
for _, r in top_neg.iterrows():
    bullet(f"{r['Aspect']}: {int(r['Negative'])} negative vs {int(r['Positive'])} positive mentions.")

doc.add_heading("5. Recommendations for improvement", 1)
if top_neg.empty:
    para("Maintain current quality and keep monitoring reviews monthly.")
for i, (_, r) in enumerate(top_neg.iterrows(), 1):
    bullet(f"{r['Aspect']} ({int(r['Negative'])} negative mentions): {RECO.get(r['Aspect'], 'Investigate the underlying complaints and address them with the vendor.')}")
if not top_pos.empty:
    bullet("Protect and promote strengths: feature " + ", ".join(top_pos['Aspect'].str.lower()) + " prominently in product-page copy and marketing.")
bullet("Reply publicly to low-rated reviews and track sentiment monthly to measure whether changes work.")

doc.add_heading("6. Limitations", 1)
for t in ("Only reviews visible on the public page were collected; the sample may not represent all customers.",
          "Aspect detection is keyword-based, and VADER is rule-based - it can miss sarcasm and mixed opinions. Spot-check the flagged rows.",
          "Relative dates such as '3 months ago' are converted approximately to YYYY-MM-DD.",
          "Counts are mentions, not unique customers; one review can mention several aspects."):
    bullet(t)

doc.save(out_path)
print("saved", out_path)
